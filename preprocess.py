#!/usr/bin/env python3
"""Read XLSX checkpoint rows and normalize them with Gemini CLI. Python 3.10+."""
from __future__ import annotations

import argparse
from collections import Counter
from contextlib import contextmanager
from datetime import date, datetime, timezone
import csv
import hashlib
import html
import io
import json
import math
import os
from pathlib import Path
import re
import shutil
import signal
import sqlite3
import subprocess
import sys
import tempfile
import time
import unicodedata
import uuid

ROOT = Path(__file__).resolve().parent
SCHEMA = json.loads((ROOT / "analysis.schema.json").read_text(encoding="utf-8"))
PROMPT = (ROOT / "prompt.txt").read_text(encoding="utf-8")
VERSION = "1.0.0"
FIELDS = {
    "checkpoint": "checkpoint", "specification": "specification",
    "설계치": "design_template", "참조문서": "reference_document",
    "pin_info": "pin_info", "adminno": "adminno", "seqno": "seqno"
}


class FormatError(ValueError):
    pass


class CliError(RuntimeError):
    pass


def dumps(value):
    return json.dumps(value, ensure_ascii=False, sort_keys=True, allow_nan=False)


def make_stdin(payload):
    # Headless Gemini CLI preprocesses @file before the model sees the prompt.
    # JSON unicode escaping is reversible and leaves source/evidence text unchanged.
    return PROMPT.replace("@", "골뱅이 기호") + "\n\n" + dumps(payload).replace("@", r"\u0040")


def normalize_text(value):
    if value is None:
        return ""
    value = unicodedata.normalize("NFC", html.unescape(str(value)))
    value = value.replace("\r\n", "\n").replace("\r", "\n").replace("\u00a0", " ")
    # Decode only presentation escapes. Do not delete /OE, backslashes in paths, or pipes.
    value = re.sub(r"\\([_.\[\]])", r"\1", value)
    return "\n".join(re.sub(r"[ \t]+", " ", line).strip() for line in value.split("\n")).strip()


def cell_value(value):
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, float):
        if not math.isfinite(value):
            return str(value)
        if value.is_integer():
            return int(value)
    return value if value is None or isinstance(value, (str, int, float, bool)) else str(value)


def header_key(value):
    return normalize_text(value).casefold().replace(" ", "")


def build_records(rows, sheet, header_row):
    iterator = iter(rows)
    try:
        headers = [str(v).strip() if v is not None else "" for v in next(iterator)]
    except StopIteration:
        raise ValueError("헤더 행이 없습니다")
    keys = [header_key(h) for h in headers]
    positions = {}
    for column, key in enumerate(keys, 1):
        if key:
            positions.setdefault(key, []).append(column)
    duplicates = {key: cols for key, cols in positions.items() if len(cols) > 1}
    used_duplicates = {key: cols for key, cols in duplicates.items() if key in FIELDS}
    if used_duplicates:
        detail = "; ".join(
            f"{headers[cols[0] - 1]!r}: " + ", ".join(f"{col}열({headers[col - 1]!r})" for col in cols)
            for key, cols in used_duplicates.items())
        raise ValueError(f"처리에 사용하는 헤더가 정규화 후 중복됩니다: {detail}. "
                         "실제로 사용할 열 하나만 해당 이름으로 두고, 나머지는 다른 이름으로 바꾸세요.")
    for required in ("checkpoint", "specification", "설계치"):
        if keys.count(required) != 1:
            raise ValueError(f"필수 열 {required!r}이 없거나 중복입니다. 실제 헤더: {headers}")
    source_headers, renamed = list(headers), []
    reserved = set(headers)
    for column, (original, key) in enumerate(zip(headers, keys), 1):
        if key not in duplicates:
            continue
        stored_as = f"{original} [col {column}]"
        suffix = 2
        while stored_as in reserved:
            stored_as = f"{original} [col {column}.{suffix}]"
            suffix += 1
        source_headers[column - 1] = stored_as
        reserved.add(stored_as)
        renamed.append({"column": column, "original": original, "stored_as": stored_as})
    records = []
    for excel_row, values in enumerate(iterator, header_row + 1):
        if not any(v is not None and str(v).strip() for v in values):
            continue
        source = {h: cell_value(v) for h, v in zip(source_headers, values) if h}
        normalized = {FIELDS[k]: normalize_text(source[h]) for h, k in zip(source_headers, keys) if k in FIELDS}
        normalized = {k: normalized.get(k, "") for k in ("checkpoint", "specification", "design_template", "reference_document", "pin_info", "adminno", "seqno")}
        identifiers = (normalized["adminno"], normalized["seqno"])
        checkpoint_id = ":".join(identifiers) if all(identifiers) else f"row:{excel_row}"
        formula_fields = [h for h, k in zip(source_headers, keys) if k in FIELDS and isinstance(source[h], str) and source[h].startswith("=")]
        records.append({
            "row_uid": f"{sheet}:{excel_row}", "checkpoint_id": checkpoint_id,
            "sheet": sheet, "excel_row": excel_row,
            "source_columns": source, "normalized_source": normalized,
            **({"renamed_headers": renamed} if renamed else {}),
            "source_issues": (["formula_requires_cached_value"] if formula_fields else []) +
                             (["missing_checkpoint_identifiers"] if not all(identifiers) else []),
            "formula_fields": formula_fields,
            "eligible": bool(normalized["checkpoint"] or normalized["specification"]) and not formula_fields
        })
    counts = Counter(r["checkpoint_id"] for r in records)
    for row in records:
        if counts[row["checkpoint_id"]] > 1:
            row["source_issues"].append("duplicate_checkpoint_id")
    return records


def read_xlsx(path, sheet, header_row):
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise ValueError("먼저 python -m pip install -r requirements.txt 를 실행하세요") from exc
    # Read formulas as formulas: never mistake an uncalculated cached null for an empty rule.
    book = load_workbook(path, read_only=True, data_only=False)
    try:
        if sheet not in book.sheetnames:
            raise ValueError(f"시트 {sheet!r}가 없습니다. 실제 시트: {book.sheetnames}")
        return build_records(book[sheet].iter_rows(min_row=header_row, values_only=True), sheet, header_row)
    finally:
        book.close()


def decode_table(path, encoding="auto"):
    data = path.read_bytes()
    encodings = (["utf-16"] if data.startswith((b"\xff\xfe", b"\xfe\xff")) else ["utf-8-sig", "cp949"]) if encoding == "auto" else [encoding]
    for candidate in encodings:
        try:
            text = data.decode(candidate)
        except UnicodeDecodeError:
            continue
        except LookupError as exc:
            raise ValueError(f"알 수 없는 인코딩: {candidate}") from exc
        if "\x00" in text:
            raise ValueError("TXT/CSV에 NUL 문자가 있습니다. 실제 텍스트 변환 여부와 --encoding utf-16-le 등을 확인하세요.")
        return text.lstrip("\ufeff"), candidate
    raise ValueError(f"인코딩을 읽지 못했습니다 ({', '.join(encodings)}). --encoding utf-8 또는 --encoding cp949 등으로 지정하세요.")


def parse_delimiter(value):
    aliases = {"comma": ",", "tab": "\t", r"\t": "\t", "semicolon": ";", "pipe": "|"}
    value = aliases.get(value, value)
    if value != "auto" and (len(value) != 1 or value in ('"', "\n", "\r")):
        raise ValueError("구분자는 auto, comma, tab, semicolon, pipe 또는 하나의 문자로 지정하세요")
    return value


def read_delimited(path, header_row=1, encoding="auto", delimiter="auto"):
    text, used_encoding = decode_table(path, encoding)
    delimiter = parse_delimiter(delimiter)
    candidates = [",", "\t", ";", "|"] if delimiter == "auto" else [delimiter]
    csv.field_size_limit(max(csv.field_size_limit(), 10_000_000))
    matches = []
    for candidate in candidates:
        reader = csv.reader(io.StringIO(text, newline=""), delimiter=candidate, strict=True)
        try:
            for _ in range(header_row - 1):
                next(reader)
            headers = next(reader)
        except (StopIteration, csv.Error):
            continue
        keys = [header_key(h) for h in headers]
        if all(k in keys for k in ("checkpoint", "specification", "설계치")):
            matches.append((candidate, headers, reader))
    if len(matches) != 1:
        raise ValueError("CSV 헤더 또는 구분자를 확인하세요. Checkpoint, Specification, 설계치 열이 필요합니다. "
                         "--delimiter tab/comma 및 --header-row를 지정할 수 있습니다. "
                         "셀을 한 줄씩 나열한 TXT는 CSV 표 구조가 없어 읽을 수 없습니다.")
    used_delimiter, headers, reader = matches[0]
    values = [headers]
    try:
        for record_row, row in enumerate(reader, header_row + 1):
            if row and len(row) != len(headers):
                raise ValueError(f"CSV 레코드 {record_row}의 열 수 {len(row)}가 헤더 {len(headers)}와 다릅니다. "
                                 "셀 안의 구분자와 개행은 CSV 따옴표로 감싸야 합니다.")
            values.append(row)
    except csv.Error as exc:
        raise ValueError(f"CSV 형식 오류 (물리 줄 {reader.line_num}): {exc}. 따옴표와 셀 내부 개행을 확인하세요.") from exc
    records = build_records(values, f"csv:{path.name}", header_row)
    for record in records:
        record["record_row"] = record["excel_row"]
        record["excel_row"] = None
        record["sheet"] = None
        record["input_format"] = "csv"
        record["input_encoding"] = used_encoding
        record["input_delimiter"] = used_delimiter
    return records


def read_input(path, sheet, header_row, encoding="auto", delimiter="auto"):
    path = Path(path)
    if not path.is_file():
        raise ValueError(f"입력 파일을 찾지 못했습니다: {path.resolve()}\n현재 폴더: {Path.cwd()}\n실제 파일명 또는 전체 경로를 지정하세요.")
    suffix = path.suffix.casefold()
    if suffix in (".xlsx", ".xlsm"):
        records = read_xlsx(path, sheet, header_row)
        for record in records:
            record["input_format"] = "xlsx"
        return records
    if suffix in (".csv", ".txt", ".tsv"):
        return read_delimited(path, header_row, encoding, delimiter)
    raise ValueError(f"지원하지 않는 확장자 {suffix!r}. XLSX, CSV, TXT 또는 TSV 파일을 지정하세요.")


def validate_schema(value, schema, path="analysis"):
    """Validate the exact JSON-Schema subset used in analysis.schema.json; no dependency."""
    kinds = schema.get("type", [])
    kinds = [kinds] if isinstance(kinds, str) else kinds
    matches = {
        "object": isinstance(value, dict), "array": isinstance(value, list),
        "string": isinstance(value, str), "null": value is None,
        "number": type(value) in (int, float) and math.isfinite(value)
    }
    if kinds and not any(matches.get(k, False) for k in kinds):
        raise FormatError(f"{path}: expected {kinds}, got {type(value).__name__}")
    if "enum" in schema and value not in schema["enum"]:
        raise FormatError(f"{path}: 허용되지 않는 값 {value!r}")
    if isinstance(value, str) and len(value.strip()) < schema.get("minLength", 0):
        raise FormatError(f"{path}: 빈 문자열")
    if isinstance(value, dict):
        missing = set(schema.get("required", [])) - set(value)
        if missing:
            raise FormatError(f"{path}: 필수 필드 누락 {sorted(missing)}")
        props = schema.get("properties", {})
        if schema.get("additionalProperties") is False and set(value) - set(props):
            raise FormatError(f"{path}: 예상하지 않은 필드 {sorted(set(value) - set(props))}")
        for name, child in props.items():
            if name in value:
                validate_schema(value[name], child, f"{path}.{name}")
    if isinstance(value, list):
        if len(value) < schema.get("minItems", 0):
            raise FormatError(f"{path}: 항목 수 부족")
        for i, item in enumerate(value):
            validate_schema(item, schema.get("items", {}), f"{path}[{i}]")


def validate_analysis(value, source):
    validate_schema(value, SCHEMA)
    for check in value["checks"]:
        for evidence in check["evidence"]:
            if evidence["quote"] not in source.get(evidence["field"], ""):
                raise FormatError(f"근거 인용이 원문에 없음: {evidence['field']} / {evidence['quote']!r}")


def parse_cli_output(text):
    try:
        envelope = json.loads(text)
    except (ValueError, TypeError) as exc:
        raise FormatError("CLI stdout이 JSON 객체가 아닙니다") from exc
    if not isinstance(envelope, dict):
        raise FormatError("CLI JSON envelope가 객체가 아닙니다")
    if envelope.get("error"):
        raise CliError(dumps(envelope["error"]))
    response = envelope.get("response")
    if not isinstance(response, str):
        raise FormatError("CLI JSON에 문자열 response가 없습니다")
    response = response.strip()
    fence = re.fullmatch(r"```(?:json)?\s*\n?(.*?)\n?```", response, flags=re.DOTALL | re.IGNORECASE)
    if fence:
        response = fence.group(1).strip()
    try:
        return json.loads(response, parse_constant=lambda x: (_ for _ in ()).throw(ValueError(x)))
    except ValueError as exc:
        raise FormatError("모델 response가 유효한 JSON이 아닙니다") from exc


def validate_items(payload, sources):
    if not isinstance(payload, dict) or set(payload) != {"items"} or not isinstance(payload["items"], list):
        raise FormatError("최상위 응답은 items 배열만 갖는 객체여야 합니다")
    seen, valid, invalid = set(), {}, {}
    for item in payload["items"]:
        if not isinstance(item, dict) or set(item) != {"request_id", "analysis"}:
            raise FormatError("items 원소는 request_id와 analysis를 가져야 합니다")
        key = item["request_id"]
        if not isinstance(key, str) or key not in sources or key in seen:
            raise FormatError("응답 ID가 입력에 없거나 중복입니다")
        seen.add(key)
        try:
            validate_analysis(item["analysis"], sources[key])
            valid[key] = item["analysis"]
        except FormatError as exc:
            invalid[key] = str(exc)
    for key in sources.keys() - seen:
        invalid[key] = "모델 응답에서 행 누락"
    return valid, invalid


def request_key(source, fingerprint):
    return hashlib.sha256((fingerprint + "\n" + dumps(source)).encode("utf-8")).hexdigest()


class Cache:
    def __init__(self, path):
        self.db = sqlite3.connect(path)
        self.db.execute("CREATE TABLE IF NOT EXISTS results (key TEXT PRIMARY KEY, analysis TEXT NOT NULL)")
        self.db.execute("CREATE TABLE IF NOT EXISTS failures (key TEXT PRIMARY KEY, error TEXT NOT NULL)")
        self.db.commit()

    def __enter__(self):
        return self

    def __exit__(self, *args):
        self.db.close()

    def get(self, key):
        found = self.db.execute("SELECT analysis FROM results WHERE key=?", (key,)).fetchone()
        return json.loads(found[0]) if found else None

    def put(self, key, value):
        self.db.execute("INSERT OR REPLACE INTO results VALUES (?,?)", (key, dumps(value)))
        self.db.execute("DELETE FROM failures WHERE key=?", (key,))
        self.db.commit()

    def fail(self, key, error):
        self.db.execute("INSERT OR REPLACE INTO failures VALUES (?,?)", (key, str(error)))
        self.db.commit()

    def error(self, key):
        found = self.db.execute("SELECT error FROM failures WHERE key=?", (key,)).fetchone()
        return found[0] if found else None


def make_batches(items, count, chars):
    batch, size = [], 0
    for item in items:
        item_size = len(dumps({"request_id": item[0], "source": item[1]}))
        if batch and (len(batch) >= count or size + item_size > chars):
            yield batch
            batch, size = [], 0
        batch.append(item)
        size += item_size
    if batch:
        yield batch


def csv_safe(text):
    text = str(text)
    return "'" + text if text.lstrip().startswith(("=", "+", "-", "@")) else text


def atomic_write(path, content):
    temporary = path.with_name(path.name + ".tmp")
    temporary.write_text(content, encoding="utf-8")
    os.replace(temporary, path)


@contextmanager
def run_lock(folder):
    path = folder / "run.lock"
    try:
        fd = os.open(path, os.O_CREAT | os.O_EXCL | os.O_WRONLY, 0o600)
    except FileExistsError as exc:
        raise ValueError(f"다른 실행 또는 이전 비정상 종료의 잠금 파일이 있습니다: {path}. 프로세스 종료를 확인한 뒤 삭제하세요.") from exc
    os.write(fd, str(os.getpid()).encode("ascii"))
    os.close(fd)
    try:
        yield
    finally:
        path.unlink(missing_ok=True)


class Gemini:
    def __init__(self, args, folder):
        self.args, self.folder, self.last_started = args, folder, 0.0
        self.executable = shutil.which(args.gemini)
        if not self.executable:
            raise ValueError(f"Gemini CLI를 찾지 못했습니다: {args.gemini}. 설치 및 로그인 후 실행하세요.")
        self.calls = 0

    @staticmethod
    def stop(process):
        if process.poll() is not None:
            return
        if os.name == "nt":
            subprocess.run(["taskkill", "/PID", str(process.pid), "/T", "/F"], capture_output=True)
        else:
            try:
                os.killpg(process.pid, signal.SIGKILL)
            except ProcessLookupError:
                pass
        process.communicate()

    def call(self, batch):
        wait = self.args.interval - (time.monotonic() - self.last_started)
        if wait > 0:
            time.sleep(wait)
        payload = {"analysis_schema": SCHEMA, "rows": [{"request_id": key, "source": source} for key, source in batch]}
        # Fixed short argument; all spreadsheet content and the full instruction go through stdin.
        command = [self.executable, "--prompt", "Process the stdin checkpoint data using its instructions. Return only JSON.",
                   "--output-format", "json", "--approval-mode", "default", "--extensions", "none"]
        if self.args.model:
            command += ["--model", self.args.model]
        self.last_started = time.monotonic()
        self.calls += 1
        call_id = uuid.uuid4().hex
        with tempfile.TemporaryDirectory(prefix="checkpoint-gemini-") as directory:
            work = Path(directory)
            (work / ".gemini").mkdir()
            # Project settings only; leave enterprise system settings and user auth intact.
            settings = {"tools": {"core": []}, "mcp": {"allowed": []},
                        "hooks": {"enabled": False}, "context": {"fileName": "CHECKPOINT_EMPTY_CONTEXT.md"},
                        "general": {"enableAutoUpdate": False}}
            (work / ".gemini" / "settings.json").write_text(dumps(settings), encoding="utf-8")
            environment = dict(os.environ, NO_COLOR="1", PYTHONIOENCODING="utf-8")
            process = subprocess.Popen(command, cwd=work, env=environment, stdin=subprocess.PIPE,
                                       stdout=subprocess.PIPE, stderr=subprocess.PIPE, text=True,
                                       encoding="utf-8", errors="replace", start_new_session=os.name != "nt")
            try:
                stdout, stderr = process.communicate(make_stdin(payload), timeout=self.args.timeout)
            except subprocess.TimeoutExpired as exc:
                self.stop(process)
                raise CliError(f"CLI timeout ({self.args.timeout}s)") from exc
            except KeyboardInterrupt:
                self.stop(process)
                raise
        logs = self.folder / "cli_logs"
        logs.mkdir(exist_ok=True)
        (logs / f"{call_id}.json").write_text(dumps({"returncode": process.returncode, "stdout": stdout, "stderr": stderr}), encoding="utf-8")
        if process.returncode:
            raise CliError(f"CLI exit {process.returncode}: {(stderr or stdout)[-2000:]}")
        return parse_cli_output(stdout)


def process_batch(batch, cache, client, attempts, backoff):
    remaining = list(batch)
    last_invalid = {}
    for attempt in range(attempts):
        try:
            payload = client.call(remaining)
            valid, invalid = validate_items(payload, dict(remaining))
        except CliError as exc:
            # Authentication/configuration failures should not become thousands of row retries.
            text = str(exc).casefold()
            permanent = any(s in text for s in ("401", "403", "unauthorized", "authentication", "not authenticated", "invalid api key", "cli exit 42", "unknown argument", "login"))
            if permanent or attempt + 1 == attempts:
                for key, _ in remaining:
                    cache.fail(key, str(exc))
                raise
            delay = min(60.0, backoff * (2 ** attempt))
            print(f"CLI 오류, {delay:g}초 뒤 재시도 {attempt + 1}/{attempts}: {str(exc)[:180]}", flush=True)
            time.sleep(delay)
            continue
        except FormatError as exc:
            valid, invalid = {}, {key: str(exc) for key, _ in remaining}
        for key, value in valid.items():
            cache.put(key, value)
        remaining = [(key, source) for key, source in remaining if key not in valid]
        last_invalid = invalid
        if not remaining:
            return
        if attempt + 1 < attempts:
            print(f"형식/근거 검증 실패 {len(remaining)}행: 남은 행만 재시도", flush=True)
    if len(remaining) > 1:
        # Smaller batches isolate a difficult row instead of losing the whole batch.
        middle = len(remaining) // 2
        process_batch(remaining[:middle], cache, client, attempts, backoff)
        process_batch(remaining[middle:], cache, client, attempts, backoff)
    else:
        for key, _ in remaining:
            cache.fail(key, last_invalid.get(key, "응답 검증 실패"))


def export_results(rows, cache, folder, info):
    results, counts = [], Counter()
    review_path = folder / "review.csv"
    review_temp = folder / "review.csv.tmp"
    with review_temp.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(["row_uid", "checkpoint_id", "status", "review_state", "summary", "checks", "issues", "error"])
        for row in rows:
            key = row.get("request_id")
            value = cache.get(key) if key else None
            error = cache.error(key) if key and value is None else None
            status = "normalized" if value is not None else "error" if error else "pending" if row["eligible"] else "skipped"
            review_state = "needs_review" if error or row["source_issues"] or (value and value["issues"]) else "candidate" if value else "not_processed"
            output = dict(row, schema_version=VERSION, status=status, review_state=review_state, analysis=value, error=error)
            if value:
                output["analysis"] = dict(value, checks=[dict(check, requirement_id=f"{row['row_uid']}:R{i}") for i, check in enumerate(value["checks"], 1)])
            results.append(dumps(output))
            counts[status] += 1
            counts["review_" + review_state] += 1
            writer.writerow([csv_safe(v) for v in (
                row["row_uid"], row["checkpoint_id"], status, review_state,
                value["summary"] if value else "", dumps(value["checks"]) if value else "",
                dumps({"source": row["source_issues"], "analysis": value["issues"] if value else []}), error or "")])
    os.replace(review_temp, review_path)
    atomic_write(folder / "checkpoints.jsonl", "\n".join(results) + ("\n" if results else ""))
    summary = dict(info, counts=dict(counts), selected_rows=len(rows), finished_at=datetime.now(timezone.utc).isoformat())
    atomic_write(folder / "summary.json", dumps(summary) + "\n")
    return summary


def main(argv=None):
    parser = argparse.ArgumentParser(description="Gemini CLI 기반 체크포인트 XLSX/CSV/TXT 전처리 (원본 수정 없음)")
    parser.add_argument("input", type=Path)
    parser.add_argument("--sheet", default="체크포인트_REF_GUIDE", help="XLSX 시트 이름; CSV/TXT에는 적용하지 않음")
    parser.add_argument("--header-row", type=int, default=1, help="헤더 행/CSV 논리 레코드 번호")
    parser.add_argument("--encoding", default="auto", help="CSV/TXT 인코딩: auto(UTF-8, BOM UTF-16, CP949) 또는 직접 지정")
    parser.add_argument("--delimiter", default="auto", help="CSV/TXT 구분자: auto, comma, tab, semicolon, pipe 또는 한 문자")
    parser.add_argument("--output", type=Path, default=Path("preprocessed"))
    parser.add_argument("--limit", type=int, help="빈 행 제외 상위 N행만 처리")
    parser.add_argument("--batch-size", type=int, default=10)
    parser.add_argument("--max-batch-chars", type=int, default=18000, help="행 payload 문자 수 목표; 단일 긴 행은 자르지 않음")
    parser.add_argument("--model", help="Gemini CLI에서 사용 가능한 모델 ID; 생략 시 CLI 기본 모델")
    parser.add_argument("--gemini", default="gemini", help="Gemini CLI 실행 파일 경로")
    parser.add_argument("--attempts", type=int, default=3)
    parser.add_argument("--timeout", type=float, default=180)
    parser.add_argument("--interval", type=float, default=2)
    parser.add_argument("--backoff", type=float, default=10)
    parser.add_argument("--cache-tag", default="", help="기본 모델/CLI 설정 변경 시 새 태그로 캐시 분리")
    parser.add_argument("--dry-run", action="store_true", help="입력 정규화만 수행; Gemini 호출 없음")
    args = parser.parse_args(argv)
    if min(args.header_row, args.batch_size, args.max_batch_chars, args.attempts) < 1 or args.timeout <= 0 or args.interval < 0 or args.backoff < 0 or (args.limit is not None and args.limit < 1):
        parser.error("행/배치/시도 횟수/timeout/limit은 양수, interval/backoff는 0 이상이어야 합니다")
    if args.model and not re.fullmatch(r"[A-Za-z0-9._:/-]+", args.model):
        parser.error("모델 ID 형식이 올바르지 않습니다")
    args.output = args.output.resolve()
    args.output.mkdir(parents=True, exist_ok=True)
    with run_lock(args.output):
        rows = read_input(args.input, args.sheet, args.header_row, args.encoding, args.delimiter)
        total = len(rows)
        if args.limit:
            rows = rows[:args.limit]
        fingerprint = hashlib.sha256(dumps({"version": VERSION, "prompt": PROMPT, "schema": SCHEMA,
            "model": args.model or "cli-default", "cache_tag": args.cache_tag}).encode("utf-8")).hexdigest()
        sources = {}
        for row in rows:
            if row["eligible"]:
                # Metadata is preserved locally but not sent to Gemini (checker IDs/dates etc.).
                source = {k: v for k, v in row["normalized_source"].items() if k not in ("adminno", "seqno")}
                key = request_key(source, fingerprint)
                row["request_id"] = key
                sources[key] = source
        atomic_write(args.output / "prepared.jsonl", "\n".join(dumps(r) for r in rows) + ("\n" if rows else ""))
        input_format = "xlsx" if args.input.suffix.casefold() in (".xlsx", ".xlsm") else "csv"
        info = {"input": str(args.input.resolve()), "sheet": args.sheet if input_format == "xlsx" else None,
                "input_format": input_format, "input_encoding": rows[0].get("input_encoding") if rows else None,
                "input_delimiter": rows[0].get("input_delimiter") if rows else None, "total_nonempty_rows": total,
                "fingerprint": fingerprint, "requested_model": args.model or "cli-default", "dry_run": args.dry_run}
        with Cache(args.output / "state.sqlite3") as cache:
            pending = [(key, source) for key, source in sources.items() if cache.get(key) is None]
            info.update(unique_requests=len(sources), cached_requests=len(sources) - len(pending))
            print(f"행 {len(rows)}/{total}, 고유 요청 {len(sources)}, 캐시 {len(sources)-len(pending)}, 남은 요청 {len(pending)}", flush=True)
            outcome, client = 0, None
            try:
                if pending and not args.dry_run:
                    client = Gemini(args, args.output)
                    batches = list(make_batches(pending, args.batch_size, args.max_batch_chars))
                    for i, batch in enumerate(batches, 1):
                        print(f"배치 {i}/{len(batches)} ({len(batch)}행) 처리", flush=True)
                        process_batch(batch, cache, client, args.attempts, args.backoff)
            except KeyboardInterrupt:
                print("중단됨: 성공 결과를 보존하고 출력 파일을 갱신합니다.", flush=True)
                outcome = 130
            except (CliError, OSError, ValueError) as exc:
                print(f"실행 중단: {exc}", file=sys.stderr)
                info["run_error"] = str(exc)
                outcome = 2
            finally:
                info["cli_calls_this_run"] = client.calls if client else 0
                summary = export_results(rows, cache, args.output, info)
                print(dumps(summary["counts"]), flush=True)
            if not args.dry_run and any(summary["counts"].get(k, 0) for k in ("pending", "error")):
                outcome = outcome or 2
            return outcome


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except (ValueError, OSError) as exc:
        print(f"오류: {exc}", file=sys.stderr)
        raise SystemExit(2)
